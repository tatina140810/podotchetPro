"""Excel для отчёта «Цепочка расходов» — 4 листа:

1. «Дерево» — человек → получено / расходы по категориям → операции / передано,
   дети вложены; группировка строк Excel (+/−). Суммы разделов — формулы SUM.
2. «По категориям» — конечные расходы всей цепочки по статьям (без передач).
3. «Все операции» — плоский список с колонкой «Цепочка» для фильтра.
4. «Сверка» — формула баланса по каждому человеку и по цепочке.
Вход — уже сконвертированный в валюту отчёта JSON из routers/reports_chain.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_SYM = {"KGS": "с", "USD": "$", "EUR": "€", "RUB": "₽"}
_HEAD_FILL = PatternFill("solid", fgColor="4F46E5")
_DEPTH_FILLS = ["C7D2FE", "DDE3FD", "EEF1FE", "F5F7FF"]
_SECTION_FILL = PatternFill("solid", fgColor="F3F4F6")
_BAD_FILL = PatternFill("solid", fgColor="FECACA")
_MONEY = '#,##0.00;[Red]-#,##0.00'
_MAX_OUTLINE = 7  # предел Excel


def _header(ws, titles, widths):
    ws.append(titles)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _marks(r: dict) -> str:
    out = []
    if r.get("status") == "pending":
        out.append("ждёт проверки")
    if r.get("report_hold"):
        out.append("ждёт выпуска")
    if r.get("auto"):
        out.append("списано при выдаче")
    if r.get("receipts"):
        out.append(f"чеков: {r['receipts']}")
    return ", ".join(out)


def _paths(nodes: list[dict]) -> dict[int, str]:
    by_id = {n["user_id"]: n for n in nodes}
    paths: dict[int, str] = {}
    for n in nodes:
        chain, cur = [], n
        while cur is not None:
            chain.append(cur["name"])
            cur = by_id.get(cur["parent_id"])
        paths[n["user_id"]] = " › ".join(reversed(chain))
    return paths


# ---------- Лист 1: дерево ----------
# Колонки: A Название, B Дата, C Тип, D На начало, E Получено, F Потрачено,
# G Передано, H На конец, I Сумма операции/раздела, J Исх. сумма, K Валюта,
# L Описание, M Отметки
def _tree_sheet(wb: Workbook, data: dict, sym: str):
    ws = wb.active
    ws.title = "Дерево"
    _header(ws, ["Кто / раздел / операция", "Дата", "Тип", f"На начало, {sym}",
                 f"Получено, {sym}", f"Потрачено, {sym}", f"Передано, {sym}",
                 f"На конец, {sym}", f"Сумма, {sym}", "Сумма в валюте", "Валюта",
                 "Описание", "Отметки"],
                [42, 11, 18, 14, 14, 14, 14, 14, 14, 13, 8, 44, 30])
    ws.sheet_properties.outlinePr.summaryBelow = False
    nodes = data["nodes"]
    children: dict = {}
    for n in nodes:
        children.setdefault(n["parent_id"], []).append(n)

    def level(r, lvl, hidden=False):
        ws.row_dimensions[r].outlineLevel = min(lvl, _MAX_OUTLINE)
        if hidden:
            ws.row_dimensions[r].hidden = True

    def op_row(r: dict, indent: str, lvl: int, sign: int):
        ws.append([indent + (r["counterparty"] or r["label"] or ""), r["date"], r["type"],
                   None, None, None, None, None, sign * r["kgs"], sign * r["amount"],
                   r["currency"], r.get("note") or "", _marks(r)])
        level(ws.max_row, lvl, hidden=True)

    def section(title: str, rows: list, indent: str, lvl: int, sign: int):
        """Заголовок раздела с формулой SUM по строкам ниже. Возвращает адрес суммы."""
        ws.append([indent + title])
        head = ws.max_row
        for c in ws[head]:
            c.fill = _SECTION_FILL
            c.font = Font(bold=True)
        level(head, lvl)
        for r in rows:
            op_row(r, indent + "    ", lvl + 1, sign)
        last = ws.max_row
        ws.cell(head, 9, f"=SUM(I{head + 1}:I{last})" if last > head else 0)
        return f"I{head}"

    def person(n: dict):
        d = n["depth"]
        indent = "    " * d
        ws.append([indent + n["name"], None, "Сотрудник", n["opening"]])
        pr = ws.max_row
        fill = PatternFill("solid", fgColor=_DEPTH_FILLS[min(d, len(_DEPTH_FILLS) - 1)])
        for c in ws[pr]:
            c.fill = fill
            c.font = Font(bold=True)
        level(pr, d)
        sub = indent + "    "
        got = [section("Получено внутри цепочки", n["incoming_internal"], sub, d + 1, 1),
               section("Получено извне", n["incoming_external"], sub, d + 1, 1)]
        # Расходы: заголовок → категории (SUM операций) → операции.
        ws.append([sub + "Расходы по категориям"])
        exp_head = ws.max_row
        for c in ws[exp_head]:
            c.fill = _SECTION_FILL
            c.font = Font(bold=True)
        level(exp_head, d + 1)
        cat_cells = [section(c["name"], c["rows"], sub + "    ", d + 2, -1) for c in n["categories"]]
        for cell in cat_cells:
            ws.row_dimensions[int(cell[1:])].hidden = True
        ws.cell(exp_head, 9, "=" + "+".join(cat_cells) if cat_cells else 0)
        gave = [section("Передано внутри цепочки", n["outgoing_internal"], sub, d + 1, -1),
                section("Ушло за пределы цепочки", n["outgoing_external"], sub, d + 1, -1)]
        ws.cell(pr, 5, f"={got[0]}+{got[1]}")
        ws.cell(pr, 6, f"=-I{exp_head}")
        ws.cell(pr, 7, f"=-({gave[0]}+{gave[1]})")
        ws.cell(pr, 8, n["closing"])
        if not n["check_ok"]:
            ws.cell(pr, 13, f"НЕ СХОДИТСЯ на {n['check_diff']}").fill = _BAD_FILL
        for kid in children.get(n["user_id"], []):
            person(kid)

    for root in children.get(None, []):
        person(root)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=10):
        for c in row:
            c.number_format = _MONEY


# ---------- Лист 2: по категориям ----------
def _categories_sheet(wb: Workbook, data: dict, sym: str):
    ws = wb.create_sheet("По категориям")
    _header(ws, ["Категория", "Операций", f"Сумма, {sym}"], [44, 11, 16])
    for c in data["categories"]:
        ws.append([c["name"], c["count"], c["total"]])
    last = ws.max_row
    ws.append(["ИТОГО конечных расходов цепочки", f"=SUM(B2:B{last})", f"=SUM(C2:C{last})"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = _MONEY


# ---------- Лист 3: все операции ----------
def _operations_sheet(wb: Workbook, data: dict, sym: str):
    ws = wb.create_sheet("Все операции")
    _header(ws, ["Цепочка", "Сотрудник", "Дата", "Тип", "Кому / от кого / категория",
                 f"Сумма, {sym}", "Сумма в валюте", "Валюта", "Движение", "Описание",
                 "Отметки"], [34, 16, 11, 18, 30, 14, 13, 8, 18, 44, 30])
    paths = _paths(data["nodes"])
    groups = [("incoming_internal", 1, "внутри цепочки"), ("incoming_external", 1, "приход извне"),
              ("outgoing_internal", -1, "внутри цепочки"), ("outgoing_external", -1, "ушло наружу")]
    rows = []
    for n in data["nodes"]:
        for key, sign, move in groups:
            rows += [(n, r, sign, move) for r in n[key]]
        for c in n["categories"]:
            rows += [(n, r, -1, "конечный расход") for r in c["rows"]]
    rows.sort(key=lambda t: (t[1]["date"], paths[t[0]["user_id"]]))
    for n, r, sign, move in rows:
        ws.append([paths[n["user_id"]], n["name"], r["date"], r["type"],
                   r["counterparty"] or r["label"], sign * r["kgs"], sign * r["amount"],
                   r["currency"], move, r.get("note") or "", _marks(r)])
    ws.auto_filter.ref = f"A1:K{max(ws.max_row, 2)}"
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for c in row:
            c.number_format = _MONEY


# ---------- Лист 4: сверка ----------
def _check_sheet(wb: Workbook, data: dict, sym: str):
    ws = wb.create_sheet("Сверка")
    _header(ws, ["Сотрудник", "На начало", "Получено внутри", "Получено извне", "Потрачено",
                 "Передано внутри", "Ушло наружу", "На конец (баланс)", "Расчёт",
                 "Разница", "Сходится"], [34, 14, 15, 15, 14, 15, 14, 16, 14, 11, 10])
    paths = _paths(data["nodes"])
    for n in data["nodes"]:
        ws.append([paths[n["user_id"]], n["opening"], n["in_internal"], n["in_external"],
                   n["expenses"], n["out_internal"], n["out_external"], n["closing"]])
        r = ws.max_row
        ws.cell(r, 9, f"=B{r}+C{r}+D{r}-E{r}-F{r}-G{r}")
        ws.cell(r, 10, f"=ROUND(I{r}-H{r},2)")
        ws.cell(r, 11, f'=IF(ABS(J{r})<=0.05,"да","НЕТ")')
    last = ws.max_row
    ws.append(["ИТОГО по цепочке"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{last})"
                                      for c in range(2, 11)])
    t = ws.max_row
    for c in ws[t]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append(["Внутренние передачи взаимно гасятся (получено внутри − передано внутри)",
               None, f"=C{t}-F{t}"])
    ws.append(["Цепочка: начало + извне − расходы − наружу − конец", None,
               f"=ROUND(B{t}+D{t}-E{t}-G{t}-H{t},2)"])
    ws.append([f"Суммы в {sym}; «на начало/конец» — тот же остаток, что в карточке сотрудника."])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=10):
        for c in row:
            c.number_format = _MONEY


def build_chain_workbook(data: dict) -> Workbook:
    sym = _SYM.get(data.get("currency", "KGS"), "с")
    wb = Workbook()
    _tree_sheet(wb, data, sym)
    _categories_sheet(wb, data, sym)
    _operations_sheet(wb, data, sym)
    _check_sheet(wb, data, sym)
    return wb

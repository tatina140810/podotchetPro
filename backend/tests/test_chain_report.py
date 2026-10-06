"""Отчёт «Цепочка расходов» (Ice → Индира → её подотчётные).

Главное — верность цифр:
- выдача внутри цепочки — не расход; расход считается один раз, у потратившего;
- «на конец» каждого узла = баланс из карточки (compute_current_balance);
- сверка узла и цепочки сходится; внутренние передачи гасятся;
- выдача с категорией (авто-расход) не задваивается;
- возврат вверх по цепочке — внутреннее движение;
- человек, получивший от двоих, в дереве один раз;
- период: остаток на начало учитывает операции до периода;
- отклонённые и удалённые расходы не считаются, неподтверждённые — с пометкой.
Плюс права и Excel.
"""
import io
from datetime import datetime
from decimal import Decimal
from types import SimpleNamespace

import openpyxl
import pytest

from conftest import auth_headers, reset_all
from models import (
    BalanceTopUp, Category, ExchangeRate, Expense, Income, MoneyTransfer, Organization, User,
)
from services.balance import compute_current_balance

AUG = datetime(2026, 8, 10)
JUL = datetime(2026, 7, 10)


@pytest.fixture()
def d(db):
    reset_all()
    org = Organization(name="Исаадат", plan="legacy")
    db.add(org)
    db.flush()
    for cur, rate in (("USD", "87.5"), ("RUB", "1.1"), ("EUR", "95")):
        db.add(ExchangeRate(org_id=org.id, from_currency=cur, to_currency="KGS", rate=Decimal(rate)))

    def user(name, role="accountable", **kw):
        u = User(org_id=org.id, name=name, phone=f"+7{name}", password_hash="x", role=role, **kw)
        db.add(u)
        db.flush()
        return u

    ns = SimpleNamespace(org=org, db=db)
    ns.ice = user("Ice", role="admin")
    ns.indira = user("Индира", role="admin")
    ns.aidina = user("Айдина")
    ns.alexa = user("Алекса")
    ns.boss = user("Босс", role="superadmin")
    ns.cat = Category(org_id=org.id, name="Хоз. расходы")
    ns.food = Category(org_id=org.id, name="Обеды")
    db.add_all([ns.cat, ns.food])
    db.flush()
    return ns


def _income(d, uid, amount, when=AUG, currency="KGS"):
    d.db.add(Income(org_id=d.org.id, amount=amount, currency=currency, source="Обменка",
                    received_by_id=uid, created_by_id=uid, date=when))


def _topup(d, frm, to, amount, when=AUG, category=None):
    t = BalanceTopUp(org_id=d.org.id, admin_id=frm, user_id=to, amount=amount, currency="KGS",
                     amount_kgs=amount, date=when, category_id=category)
    d.db.add(t)
    d.db.flush()
    return t


def _expense(d, uid, amount, when=AUG, currency="KGS", status="approved", cat=None, **kw):
    e = Expense(org_id=d.org.id, employee_id=uid, amount=amount, currency=currency,
                status=status, spent_at=when, category_id=(cat or d.cat).id, **kw)
    d.db.add(e)
    d.db.flush()
    return e


def _get(client, d, root, viewer=None, **params):
    q = {"root_id": root.id, "from": "2026-08-01", "to": "2026-08-31", **params}
    r = client.get("/api/reports/chain", headers=auth_headers(viewer or d.boss), params=q)
    assert r.status_code == 200, r.text
    data = r.json()
    data["by"] = {n["name"]: n for n in data["nodes"]}
    return data


def _base_chain(d):
    """Ice: приход 1000 → Индире 600 → Айдине 200. Расходы: Ice 100, Индира 300, Айдина 150."""
    _income(d, d.ice.id, 1000)
    _topup(d, d.ice.id, d.indira.id, 600)
    _topup(d, d.indira.id, d.aidina.id, 200)
    _expense(d, d.ice.id, 100)
    _expense(d, d.indira.id, 300)
    _expense(d, d.aidina.id, 150, cat=d.food)
    d.db.commit()


def test_three_levels_numbers_and_check(client, d):
    _base_chain(d)
    data = _get(client, d, d.ice)
    ice, ind, ai = data["by"]["Ice"], data["by"]["Индира"], data["by"]["Айдина"]
    assert [n["name"] for n in data["nodes"]] == ["Ice", "Индира", "Айдина"]
    assert ind["parent_id"] == d.ice.id and ai["parent_id"] == d.indira.id
    assert (ice["in_external"], ice["expenses"], ice["out_internal"], ice["closing"]) == (1000, 100, 600, 300)
    assert (ind["in_internal"], ind["expenses"], ind["out_internal"], ind["closing"]) == (600, 300, 200, 100)
    assert (ai["in_internal"], ai["expenses"], ai["closing"]) == (200, 150, 50)
    t = data["totals"]
    # Внутренние выдачи (800) в расходы не попали: расходы цепочки = 100+300+150.
    assert t["expenses"] == 550 and t["in_external"] == 1000 and t["closing"] == 450
    assert t["in_internal"] == t["out_internal"] == 800
    assert t["check_ok"] and all(n["check_ok"] for n in data["nodes"])
    assert {c["name"]: c["total"] for c in data["categories"]} == {"Хоз. расходы": 400, "Обеды": 150}


def test_closing_equals_card_balance(client, d):
    _base_chain(d)
    data = _get(client, d, d.ice, to="2026-12-31")
    for u in (d.ice, d.indira, d.aidina):
        d.db.expire_all()
        assert Decimal(str(data["by"][u.name]["closing"])) == compute_current_balance(d.db, d.org.id, u.id)


def test_received_from_two_appears_once(client, d):
    _base_chain(d)
    _topup(d, d.ice.id, d.aidina.id, 50)
    d.db.commit()
    data = _get(client, d, d.ice)
    assert [n["name"] for n in data["nodes"]].count("Айдина") == 1
    ai = data["by"]["Айдина"]
    assert ai["in_internal"] == 250 and ai["closing"] == 100
    assert {r["counterparty"] for r in ai["incoming_internal"]} == {"Ice", "Индира"}
    assert data["totals"]["check_ok"]


def test_return_upwards_is_internal(client, d):
    _base_chain(d)
    d.db.add(MoneyTransfer(org_id=d.org.id, from_user_id=d.aidina.id, to_user_id=d.indira.id,
                           amount=30, currency="KGS", amount_kgs=30, created_at=AUG))
    d.db.commit()
    data = _get(client, d, d.ice)
    assert data["totals"]["expenses"] == 550  # возврат не расход
    assert data["by"]["Индира"]["in_internal"] == 630
    assert data["by"]["Айдина"]["out_internal"] == 30
    assert data["totals"]["check_ok"]


def test_topup_with_category_not_doubled(client, d):
    """Выдача с категорией сразу становится расходом получателя (авто-расход)."""
    _income(d, d.ice.id, 1000)
    t = _topup(d, d.ice.id, d.alexa.id, 250, category=d.cat.id)
    _expense(d, d.alexa.id, 250, source_topup_id=t.id)
    d.db.commit()
    data = _get(client, d, d.ice)
    al = data["by"]["Алекса"]
    assert al["expenses"] == 250 and al["closing"] == 0
    assert al["categories"][0]["rows"][0]["auto"] is True
    assert data["totals"]["expenses"] == 250 and data["totals"]["check_ok"]


def test_opening_balance_from_before_period(client, d):
    _income(d, d.ice.id, 500, when=JUL)
    _topup(d, d.ice.id, d.indira.id, 200, when=JUL)
    _base_chain(d)
    data = _get(client, d, d.ice)
    assert data["by"]["Ice"]["opening"] == 300 and data["by"]["Индира"]["opening"] == 200
    assert data["by"]["Индира"]["closing"] == 300
    assert data["totals"]["check_ok"]


def test_statuses_and_currency(client, d):
    _base_chain(d)
    _expense(d, d.indira.id, 10, status="pending")
    _expense(d, d.indira.id, 999, status="rejected")
    gone = _expense(d, d.indira.id, 777)
    gone.deleted_at = datetime(2026, 8, 20)
    _expense(d, d.aidina.id, 1, currency="USD")  # 87.5 с по курсу
    d.db.commit()
    data = _get(client, d, d.ice)
    ind = data["by"]["Индира"]
    assert ind["expenses"] == 310 and ind["pending"] == 10
    assert data["by"]["Айдина"]["expenses"] == 237.5
    assert data["totals"]["check_ok"]
    usd = _get(client, d, d.ice, currency="USD")
    assert usd["currency"] == "USD" and abs(usd["totals"]["expenses"] - 647.5 / 87.5) < 0.02


def test_hidden_confidential_stays_outside(client, d):
    _base_chain(d)
    d.aidina.is_confidential = True
    d.db.commit()
    data = _get(client, d, d.ice, viewer=d.ice)  # admin не видит конфиденциальных
    assert "Айдина" not in data["by"]
    out = data["by"]["Индира"]["outgoing_external"]
    assert out[0]["counterparty"] == "Скрытый сотрудник" and out[0]["kgs"] == 200
    assert data["totals"]["out_external"] == 200 and data["totals"]["check_ok"]


def test_accountable_own_or_subordinate_chain(client, d):
    _base_chain(d)
    d.alexa.supervisor_id = d.aidina.id
    d.db.commit()

    def status_for(root):
        return client.get("/api/reports/chain", headers=auth_headers(d.aidina),
                          params={"root_id": root.id}).status_code
    assert status_for(d.ice) == 403        # чужая цепочка
    assert status_for(d.aidina) == 200     # своя
    assert status_for(d.alexa) == 200      # подчинённой


def test_excel_four_sheets_with_outline(client, d):
    _base_chain(d)
    r = client.get("/api/reports/chain.xlsx", headers=auth_headers(d.boss),
                   params={"root_id": d.ice.id, "from": "2026-08-01", "to": "2026-08-31"})
    assert r.status_code == 200, r.text
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Дерево", "По категориям", "Все операции", "Сверка"]
    tree = wb["Дерево"]
    assert max(tree.row_dimensions[i].outlineLevel for i in range(2, tree.max_row + 1)) >= 3
    check = wb["Сверка"]
    assert [check.cell(i, 1).value for i in (2, 3, 4)] == ["Ice", "Ice › Индира", "Ice › Индира › Айдина"]
    assert check.cell(5, 1).value == "ИТОГО по цепочке"


def test_workspace_topup_details_masked_for_admin(client, d):
    """Детали выдачи в проектное пространство (категория, комментарий) admin не видит,
    но сумма остаётся — сверка сходится."""
    from models import ProjectWorkspace
    ws = ProjectWorkspace(org_id=d.org.id, name="Стройка", owner_id=d.alexa.id, is_active=True)
    d.db.add(ws)
    d.db.flush()
    _income(d, d.ice.id, 1000)
    t = _topup(d, d.ice.id, d.alexa.id, 300, category=d.cat.id)
    t.workspace_id = ws.id
    t.note = "секретный объект"
    d.db.commit()
    data = _get(client, d, d.ice, viewer=d.ice)
    row = data["by"]["Ice"]["outgoing_internal"][0]
    assert row["label"] == "Проектное пространство" and row["note"] is None
    assert row["kgs"] == 300 and data["totals"]["check_ok"]
    # superadmin видит детали
    full = _get(client, d, d.ice)["by"]["Ice"]["outgoing_internal"][0]
    assert full["note"] == "секретный объект"


def test_workspace_owner_isolated(client, d):
    """Владелец пространства (даже admin) строит цепочку только по участникам."""
    from models import ProjectWorkspace, ProjectWorkspaceMember
    ws = ProjectWorkspace(org_id=d.org.id, name="Стройка", owner_id=d.indira.id, is_active=True)
    d.db.add(ws)
    d.db.flush()
    d.db.add_all([ProjectWorkspaceMember(workspace_id=ws.id, user_id=d.indira.id),
                  ProjectWorkspaceMember(workspace_id=ws.id, user_id=d.alexa.id)])
    _base_chain(d)
    _topup(d, d.indira.id, d.alexa.id, 40)
    d.db.commit()
    r = client.get("/api/reports/chain", headers=auth_headers(d.indira), params={"root_id": d.ice.id})
    assert r.status_code == 404  # Ice — не участник
    data = _get(client, d, d.indira, viewer=d.indira)
    assert set(data["by"]) == {"Индира", "Алекса"}  # Айдина вне пространства → наружу
    assert data["by"]["Индира"]["out_external"] == 200 and data["totals"]["check_ok"]

"""Excel-выгрузка «Последних операций» (экран «Импорт истории»).

Ключевое требование: выгрузка обязана содержать ровно те же операции, что
показывает список на экране при тех же фильтрах — одна формула на оба пути
(_recent_operations_rows). Проверяем: соответствие JSON-выдаче, работу фильтров,
подытоги по валютам, права (не-admin не выгружает).

Сессию сида закрываем сразу, наружу отдаём только числовые id — иначе открытая
сессия держит блокировку и TRUNCATE следующего теста виснет.
"""
import io
from types import SimpleNamespace

from openpyxl import load_workbook

import auth
import database
from conftest import reset_all
from models import User
from tests.seed import seed_all

EXPORT_URL = "/api/admin/recent-operations/export.xlsx"
LIST_URL = "/api/admin/recent-operations"


def _fresh_seed() -> SimpleNamespace:
    reset_all()
    db = database.SessionLocal()
    try:
        s = seed_all(db)
        ids = SimpleNamespace(
            org_id=s.org.id,
            director=s.director.id, acc1=s.acc1.id, acc2=s.acc2.id,
            auditor=s.auditor.id,
            depA=s.depA.id, depB=s.depB.id,
            catNormal=s.catNormal.id,
        )
    finally:
        db.close()
    return ids


def _hdr(user_id: int, org_id: int, role: str) -> dict:
    return {"Authorization": f"Bearer {auth.create_access_token(user_id, org_id, role)}"}


def _sheet(resp):
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]
    return load_workbook(io.BytesIO(resp.content)).active


def _data_rows(ws) -> list[list]:
    """Строки операций: всё до пустой строки-разделителя перед подытогами."""
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None and row[1] is None and row[2] is None:
            break
        out.append(list(row))
    return out


def test_export_matches_list_without_filters(client):
    ids = _fresh_seed()
    h = _hdr(ids.director, ids.org_id, "superadmin")

    total = client.get(f"{LIST_URL}?limit=1", headers=h).json()["total"]
    assert total > 0

    ws = _sheet(client.get(EXPORT_URL, headers=h))
    assert [c.value for c in ws[1]] == [
        "Дата", "Тип", "Кто / Кому", "Сумма", "Валюта", "Категория / Источник", "Комментарий"
    ]
    assert len(_data_rows(ws)) == total


def test_export_respects_filters(client):
    ids = _fresh_seed()
    h = _hdr(ids.director, ids.org_id, "superadmin")

    for qs in (
        f"kind=expense",
        f"employee_id={ids.acc1}",
        f"department_id={ids.depB}",
        f"category_id={ids.catNormal}",
        f"amount_min=5000",
    ):
        total = client.get(f"{LIST_URL}?limit=1&{qs}", headers=h).json()["total"]
        rows = _data_rows(_sheet(client.get(f"{EXPORT_URL}?{qs}", headers=h)))
        assert len(rows) == total, f"фильтр {qs}: в файле {len(rows)}, в списке {total}"

    # kind=expense — в файле только расходы
    rows = _data_rows(_sheet(client.get(f"{EXPORT_URL}?kind=expense", headers=h)))
    assert {r[1] for r in rows} == {"Расход"}


def test_export_totals_per_currency(client):
    ids = _fresh_seed()
    h = _hdr(ids.director, ids.org_id, "superadmin")

    ws = _sheet(client.get(f"{EXPORT_URL}?kind=expense", headers=h))
    rows = _data_rows(ws)
    expected: dict = {}
    for r in rows:
        expected[r[4]] = expected.get(r[4], 0.0) + float(r[3])

    totals = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[2] == "ИТОГО":
            totals[row[4]] = float(row[3])
    assert totals, "нет строк ИТОГО"
    assert set(totals) == set(expected)
    for cur, val in expected.items():
        assert abs(totals[cur] - val) < 0.01
    # Единого SUM по всем валютам быть не должно (валюты складывать нельзя)
    assert len(totals) == len(expected)


def _set_role(user_id: int, role: str) -> None:
    db = database.SessionLocal()
    try:
        u = db.get(User, user_id)
        u.role = role
        db.commit()
    finally:
        db.close()


def test_export_hides_confidential_for_admin(client):
    """Конфиденциальный сотрудник скрыт от admin и в выгрузке — как и в списке."""
    ids = _fresh_seed()
    _set_role(ids.acc1, "admin")  # в seed роли admin нет, а права читаются из БД
    h = _hdr(ids.acc1, ids.org_id, "admin")

    total = client.get(f"{LIST_URL}?limit=1", headers=h).json()["total"]
    rows = _data_rows(_sheet(client.get(EXPORT_URL, headers=h)))
    assert len(rows) == total
    assert not any("Билим" in str(r[2]) for r in rows)


def test_export_forbidden_for_accountable(client):
    ids = _fresh_seed()
    r = client.get(EXPORT_URL, headers=_hdr(ids.acc2, ids.org_id, "accountable"))
    assert r.status_code == 403
